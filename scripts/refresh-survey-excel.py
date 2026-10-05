#!/usr/bin/env python3
"""
Pulls every response from the BillAvya survey's Firestore collection
(surveyResponses, project letsgetdirty) and regenerates an up-to-date Excel
workbook in ~/Downloads — one row per respondent plus a per-question summary
with counts/percentages, mirroring the survey's own live results page.

Meant to be re-run on a schedule (see com.billavya.surveyexcel.plist, a
launchd agent that runs this every 2 minutes) so the file in Downloads
stays current without needing Claude or a terminal open.
"""
import json
import os
import ssl
import urllib.request
from datetime import datetime
import certifi
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ID = "letsgetdirty"
API_KEY = "AIzaSyC_-oQ0MY2anXHrD_Bm-4-exLmHRZb4Apk"
BASE_URL = f"https://firestore.googleapis.com/v1/projects/{PROJECT_ID}/databases/(default)/documents/surveyResponses?key={API_KEY}&pageSize=1000"
# ~/Downloads is a macOS privacy-protected folder — a background agent (not
# launched from Terminal, which already has that grant) gets a silent
# PermissionError writing there unless /usr/bin/python3 is explicitly added
# under System Settings -> Privacy & Security -> Full Disk Access. Until/
# unless that's granted, this writes here instead, which needs no special
# permission and updates just as live.
FALLBACK_DIR = os.path.expanduser("~/BillAvya-Survey-Live")
DOWNLOADS_PATH = os.path.expanduser("~/Downloads/BillAvya-Survey-Results.xlsx")
FALLBACK_PATH = os.path.join(FALLBACK_DIR, "BillAvya-Survey-Results.xlsx")
SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())

QUESTIONS = [
    ("monthlyOfflinePurchases", "On average, how many times a month do you shop at a physical, in-person store?",
     ["0–2", "3–5", "6–10", "More than 10"]),
    ("confidenceSharingContact", "How comfortable are you sharing your phone number or email address with a merchant solely to receive a digital invoice?",
     ["Not at all comfortable", "Slightly comfortable", "Moderately comfortable", "Comfortable", "Very comfortable"]),
    ("lostWarranty", "Has a misplaced invoice ever cost you a product's warranty coverage?", ["Yes", "No"]),
    ("troubleFindingEmailedInvoice", "Have you ever had trouble locating an invoice that was emailed to you?", ["Yes", "No"]),
    ("sharingContactIsBestWay", "In your view, is sharing your email address or phone number with merchants the best available way to receive an invoice?",
     ["Yes", "No", "Not sure"]),
    ("sharingDetailsLeadsToCybercrime", "Do you believe the widespread practice of sharing personal contact details with every merchant increases the risk of cybercrime?",
     ["Yes", "No", "Not sure"]),
    ("wantsBetterWay", "Would you welcome a better way to receive digital invoices, one that doesn't require sharing your phone number or email address?",
     ["Yes", "No"]),
    ("importanceOfOnePlace", "How important is it to you to have your spending analytics and digital invoices available in a single place?",
     ["Not important", "Slightly important", "Moderately important", "Important", "Very important"]),
    ("ageBracket", "Which age group do you belong to?", ["Under 18", "18–24", "25–34", "35–44", "45–54", "55–64", "65+"]),
    ("country", "Which country are you responding from?", ["India", "Canada", "United States", "United Kingdom", "Australia", "Other"]),
    ("financialLiteracyImportant", "Do you believe financial literacy will matter more and more in the years ahead?", ["Yes", "No", "Not sure"]),
    ("digitalInvoicingIsGreen", "Would you agree that digital invoicing is a meaningful step toward more sustainable, paperless commerce?",
     ["Yes", "No", "Not sure"]),
    ("wisdomInSpendingWisely", "Would you agree that true wisdom lies more in spending money wisely than in earning it?", ["Yes", "No", "Not sure"]),
]
NAME_FIELD = "name"
COUNTRY_OTHER_FIELD = "countryOther"

NAVY = "0B2545"
TEAL = "00C2A8"
LIGHT = "F3F6F8"
WHITE = "FFFFFF"
thin = Side(style="thin", color="E5EBED")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def fs_value(v):
    """Unwraps one Firestore REST 'Value' object to a plain Python value."""
    if "stringValue" in v:
        return v["stringValue"]
    if "integerValue" in v:
        return int(v["integerValue"])
    if "doubleValue" in v:
        return v["doubleValue"]
    if "booleanValue" in v:
        return v["booleanValue"]
    if "nullValue" in v:
        return None
    return None


def fetch_responses():
    responses = []
    page_token = None
    while True:
        url = BASE_URL + (f"&pageToken={page_token}" if page_token else "")
        with urllib.request.urlopen(url, timeout=20, context=SSL_CONTEXT) as r:
            data = json.loads(r.read().decode("utf-8"))
        for doc in data.get("documents", []):
            fields = doc.get("fields", {})
            row = {k: fs_value(v) for k, v in fields.items()}
            responses.append(row)
        page_token = data.get("nextPageToken")
        if not page_token:
            break
    return responses


def build_workbook(responses):
    wb = Workbook()

    # ---- Sheet 1: Summary (mirrors the survey's own live results view) ----
    ws = wb.active
    ws.title = "Summary"
    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = "BillAvya Customer Insights Survey — Live Summary"
    title.font = Font(name="Calibri", size=16, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:D2")
    sub = ws["A2"]
    resp_word = "response" if len(responses) == 1 else "responses"
    sub.value = f"{len(responses)} {resp_word} so far · refreshed {datetime.now().strftime('%b %d, %Y %I:%M %p')}"
    sub.font = Font(name="Calibri", size=10.5, italic=True, color="5B6B7C")
    sub.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    row = 4
    for qid, prompt, options in QUESTIONS:
        ws.merge_cells(f"A{row}:D{row}")
        qcell = ws[f"A{row}"]
        qcell.value = prompt
        qcell.font = Font(name="Calibri", size=11.5, bold=True, color=WHITE)
        qcell.fill = PatternFill("solid", fgColor=TEAL)
        qcell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 32
        row += 1

        counts = {o: 0 for o in options}
        for r in responses:
            v = r.get(qid)
            if v in counts:
                counts[v] += 1
        answered = sum(counts.values())

        headers = ["Answer", "Count", "% of answers"]
        for i, h in enumerate(headers):
            c = ws.cell(row=row, column=i + 1, value=h)
            c.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            c.border = border
        row += 1

        for o in options:
            pct = f"{(counts[o] / answered * 100):.0f}%" if answered else "0%"
            vals = [o, counts[o], pct]
            for i, v in enumerate(vals):
                c = ws.cell(row=row, column=i + 1, value=v)
                c.border = border
                c.font = Font(name="Calibri", size=10.5)
                if row % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=LIGHT)
            row += 1
        row += 1  # blank spacer row between questions

    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 10
    ws.freeze_panes = "A4"

    # ---- Sheet 2: Raw responses, one row per respondent ----
    ws2 = wb.create_sheet("Responses")
    cols = [qid for qid, _, _ in QUESTIONS] + [COUNTRY_OTHER_FIELD, NAME_FIELD, "submittedAt"]
    headers2 = [prompt for _, prompt, _ in QUESTIONS] + ["Country (other, if specified)", "Name (optional)", "Submitted At"]
    for i, h in enumerate(headers2, start=1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(name="Calibri", size=10.5, bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = border
    ws2.row_dimensions[1].height = 40

    responses_sorted = sorted(responses, key=lambda r: r.get("submittedAtMs") or 0)
    for ridx, r in enumerate(responses_sorted, start=2):
        for cidx, col in enumerate(cols, start=1):
            if col == "submittedAt":
                ms = r.get("submittedAtMs")
                val = datetime.fromtimestamp(ms / 1000).strftime("%b %d, %Y %I:%M %p") if ms else ""
            else:
                val = r.get(col, "")
            c = ws2.cell(row=ridx, column=cidx, value=val)
            c.border = border
            c.font = Font(name="Calibri", size=10)
            if ridx % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)

    for i in range(1, len(headers2) + 1):
        ws2.column_dimensions[get_column_letter(i)].width = 22
    ws2.freeze_panes = "A2"

    return wb


def main():
    responses = fetch_responses()
    wb = build_workbook(responses)
    try:
        wb.save(DOWNLOADS_PATH)
        print(f"{datetime.now().isoformat()} — saved {len(responses)} responses to {DOWNLOADS_PATH}")
    except PermissionError:
        os.makedirs(FALLBACK_DIR, exist_ok=True)
        wb.save(FALLBACK_PATH)
        print(
            f"{datetime.now().isoformat()} — saved {len(responses)} responses to {FALLBACK_PATH} "
            f"(no permission yet to write directly to Downloads — see the comment above FALLBACK_DIR)"
        )


if __name__ == "__main__":
    main()
